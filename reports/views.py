import calendar
from datetime import date
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render

from attendance.models import Attendance, Employee

MONTHS = [(i, calendar.month_name[i]) for i in range(1, 13)]


def _month_from_request(request):
    today = date.today()
    try:
        month = int(request.GET.get("month", today.month))
        year = int(request.GET.get("year", today.year))
    except (TypeError, ValueError):
        month, year = today.month, today.year
    month = min(max(month, 1), 12)
    year = min(max(year, 1), 9998)
    return month, year


def _month_bounds(month, year):
    """First day of the month and first day of the following month."""
    first_day = date(year, month, 1)
    if month == 12:
        return first_day, date(year + 1, 1, 1)
    return first_day, date(year, month + 1, 1)


def _build_month_data(month, year):
    days_in_month = calendar.monthrange(year, month)[1]
    days = list(range(1, days_in_month + 1))
    first_day, next_month = _month_bounds(month, year)

    # A date range (rather than date__year / date__month) lets the database
    # use the index on the date column.
    records = Attendance.objects.filter(
        date__gte=first_day, date__lt=next_month
    ).select_related("employee")

    by_employee = {}
    for record in records:
        by_employee.setdefault(record.employee_id, {})[record.date.day] = record

    rows = []
    for employee in Employee.objects.filter(is_active=True):
        marks = by_employee.get(employee.id, {})
        cells = []
        counts = {"P": 0, "A": 0}
        reasons = []
        for day in days:
            record = marks.get(day)
            if record:
                counts[record.status] = counts.get(record.status, 0) + 1
                if record.status in Attendance.REASON_REQUIRED and record.reason:
                    reasons.append(f"{day:02d}: {record.reason}")
                cells.append({"day": day, "code": record.status,
                              "label": record.get_status_display(),
                              "reason": record.reason})
            else:
                cells.append({"day": day, "code": "", "label": "Not marked",
                              "reason": ""})
        payable = counts["P"]
        rows.append({
            "employee": employee,
            "cells": cells,
            "counts": counts,
            "payable": payable,
            "reasons": reasons,
        })
    return days, rows


@login_required
def monthly_attendance(request):
    month, year = _month_from_request(request)
    days, rows = _build_month_data(month, year)
    return render(request, "reports/monthly_attendance.html", {
        "active": "monthly",
        "month": month,
        "year": year,
        "month_name": calendar.month_name[month],
        "months": MONTHS,
        "years": range(date.today().year - 5, date.today().year + 2),
        "days": days,
        "rows": rows,
        "total_present": sum(r["counts"]["P"] for r in rows),
        "total_absent": sum(r["counts"]["A"] for r in rows),
    })


@login_required
def monthly_attendance_xlsx(request):
    # openpyxl is slow to import, so it is only loaded when an Excel file is
    # actually requested instead of on every cold start.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    month, year = _month_from_request(request)

    wb = Workbook()
    sheet = wb.active
    sheet.title = f"{calendar.month_abbr[month]} {year}"

    header_fill = PatternFill("solid", fgColor="14532D")
    header_font = Font(color="FFFFFF", bold=True)

    sheet.append(["Date", "Status", "Reason for absent"])
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    first_day, next_month = _month_bounds(month, year)
    records = (
        Attendance.objects.filter(date__gte=first_day, date__lt=next_month)
        .select_related("employee")
        .order_by("date")
    )
    for record in records:
        is_present = record.status == Attendance.PRESENT
        sheet.append([
            record.date.strftime("%d-%m-%Y"),
            "Present" if is_present else "Absent",
            "" if is_present else record.reason,
        ])

    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 50

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    filename = f"lr-coming-{year}-{month:02d}.xlsx"
    response = HttpResponse(
        stream.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
