"""Turns a raw billing export into the six column LR pending list.

Pipeline
--------
1. read every sheet of the uploaded workbook (or CSV)
2. split single column rows into real columns (text to columns)
3. find the header row and map it onto the six columns we keep
4. drop junk rows: blanks, repeated headers, totals, page footers
5. keep only the rows whose LR number is marked OK, i.e. the bills whose
   LR number is still missing
"""

import csv
import io
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Final column order, exactly as requested.
FINAL_COLUMNS = [
    "Transport Name",
    "Party Name",
    "Bill No",
    "LR No",
    "LR Date",
    "Cases",
]

# Header spellings seen in real exports, normalised to lowercase letters only.
COLUMN_ALIASES = {
    "Transport Name": [
        "transportname", "transport", "transporter", "transportername",
        "transname", "transportnm", "trpname", "transportcompany",
    ],
    "Party Name": [
        "partyname", "party", "customer", "customername", "accountname",
        "nameofparty", "partysname", "buyer", "consignee", "consigneename",
    ],
    "Bill No": [
        "billno", "billnumber", "invoiceno", "invno", "invoicenumber",
        "billnum", "docno", "documentno", "voucherno", "vchno", "billnodot",
    ],
    "LR No": [
        "lrno", "lrnumber", "lrn", "grno", "lrgrno", "lrnodot", "lorryreceiptno",
        "lrnum", "docketno", "lrno1",
    ],
    "LR Date": [
        "lrdate", "grdate", "lrdt", "lrdatedot", "docketdate", "lrdatee",
    ],
    "Cases": [
        "cases", "case", "cs", "qty", "quantity", "noofcases", "nocases",
        "packets", "pkts", "pkt", "boxes", "bundles", "nos",
    ],
}

BILL_SERIES_ALIASES = ["billseries", "series", "billsrs", "srs", "billtype", "prefix"]

# Rows that are page furniture rather than data.
JUNK_ROW_PATTERNS = [
    r"^total", r"^grand\s*total", r"^sub\s*total", r"^page\s*no",
    r"^printed", r"^report", r"^continued", r"^opening", r"^closing",
    r"^\*+$", r"^-+$", r"^=+$",
]

DELIMITERS = ["\t", "|", "~", ";", ","]


def _norm(value):
    """Lowercase, letters and digits only - used to compare headers."""
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_blank_row(row):
    return all(not _clean(cell) for cell in row)


def _is_junk_row(row):
    joined = " ".join(_clean(cell) for cell in row).strip().lower()
    if not joined:
        return True
    for pattern in JUNK_ROW_PATTERNS:
        if re.match(pattern, joined):
            return True
    return False


# ---------------------------------------------------------------- reading


def read_rows(uploaded_file):
    """Return {sheet name: [row, row, ...]} for xlsx/xlsm/csv input."""
    name = (uploaded_file.name or "").lower()
    raw = uploaded_file.read()

    if name.endswith(".csv") or name.endswith(".txt"):
        text = raw.decode("utf-8-sig", errors="replace")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            reader = csv.reader(io.StringIO(text), dialect)
        except csv.Error:
            reader = csv.reader(io.StringIO(text))
        return {"Sheet1": [list(r) for r in reader]}

    workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    sheets = {}
    for sheet in workbook.worksheets:
        sheets[sheet.title] = [list(r) for r in sheet.iter_rows(values_only=True)]
    workbook.close()
    return sheets


# ------------------------------------------------------- text to columns


def split_text_columns(rows):
    """Excel's text to columns, done automatically.

    Some exports arrive with the whole line stuffed into column A. When that
    is the case the line is split on the delimiter that appears most often.
    """
    populated = [r for r in rows if not _is_blank_row(r)]
    if not populated:
        return rows, None

    single_column = sum(
        1 for r in populated if sum(1 for c in r if _clean(c)) <= 1
    )
    if single_column < len(populated) * 0.8:
        return rows, None      # already in proper columns

    lines = []
    for row in rows:
        cells = [c for c in row if _clean(c)]
        lines.append(_clean(cells[0]) if cells else "")

    best_delim, best_score = None, 0
    for delim in DELIMITERS:
        score = sum(1 for line in lines if line.count(delim) >= 2)
        if score > best_score:
            best_delim, best_score = delim, score

    if best_delim is None or best_score < max(3, len(lines) * 0.3):
        # Fall back to runs of two or more spaces (fixed width print outs).
        if sum(1 for line in lines if re.search(r"\s{2,}", line)) >= max(3, len(lines) * 0.3):
            return [re.split(r"\s{2,}", line) for line in lines], "multiple spaces"
        return rows, None

    return [line.split(best_delim) for line in lines], repr(best_delim)


# ------------------------------------------------------- header handling


def find_header(rows, scan_depth=25):
    """Locate the header row and map wanted columns onto their indexes."""
    best = None
    for index, row in enumerate(rows[:scan_depth]):
        mapping = {}
        series_index = None
        for col, cell in enumerate(row):
            key = _norm(cell)
            if not key:
                continue
            for target, aliases in COLUMN_ALIASES.items():
                if target in mapping:
                    continue
                if key in aliases or any(key.startswith(a) for a in aliases):
                    mapping[target] = col
                    break
            else:
                if key in BILL_SERIES_ALIASES and series_index is None:
                    series_index = col
        if len(mapping) > (len(best["mapping"]) if best else 0):
            best = {"row": index, "mapping": mapping, "series": series_index}
        if len(mapping) == len(FINAL_COLUMNS):
            break
    return best


# ------------------------------------------------------------ extraction


def is_ok_marker(value):
    """True when the LR number cell is the OK marker - the bill has no LR yet."""
    text = _clean(value).lower()
    if not text:
        return False
    return text in {"ok", "o.k", "o.k.", "ok.", "okk"} or bool(
        re.fullmatch(r"ok\W*", text)
    )


def process_workbook(uploaded_file, include_blank_lr=False, merge_series=False):
    sheets = read_rows(uploaded_file)

    report = {
        "sheets": [],
        "rows_in": 0,
        "rows_out": 0,
        "blank_rows_removed": 0,
        "junk_rows_removed": 0,
        "rows_with_lr": 0,
        "columns_dropped": 0,
        "split_note": None,
        "columns_found": [],
        "columns_missing": [],
    }
    data = []

    for sheet_name, rows in sheets.items():
        report["rows_in"] += len(rows)
        rows, split_note = split_text_columns(rows)
        if split_note and not report["split_note"]:
            report["split_note"] = split_note

        header = find_header(rows)
        if not header or len(header["mapping"]) < 2:
            report["sheets"].append({"name": sheet_name, "used": False,
                                     "reason": "no recognisable header row"})
            continue

        mapping = header["mapping"]
        series_index = header["series"] if merge_series else None
        widest = max(len(r) for r in rows) if rows else 0
        report["columns_dropped"] = max(
            report["columns_dropped"], max(widest - len(mapping), 0)
        )
        for column in mapping:
            if column not in report["columns_found"]:
                report["columns_found"].append(column)

        kept_here = 0
        for row in rows[header["row"] + 1:]:
            if _is_blank_row(row):
                report["blank_rows_removed"] += 1
                continue
            if _is_junk_row(row):
                report["junk_rows_removed"] += 1
                continue

            record = {}
            for column in FINAL_COLUMNS:
                index = mapping.get(column)
                record[column] = (
                    _clean(row[index]) if index is not None and index < len(row) else ""
                )

            # A repeated header in the middle of the print out.
            if _norm(record["Party Name"]) in COLUMN_ALIASES["Party Name"]:
                report["junk_rows_removed"] += 1
                continue
            # Nothing identifying the bill: not a data row.
            if not record["Party Name"] and not record["Bill No"] and not record["Transport Name"]:
                report["junk_rows_removed"] += 1
                continue

            if series_index is not None and series_index < len(row):
                series = _clean(row[series_index])
                if series and record["Bill No"]:
                    record["Bill No"] = f"{series}-{record['Bill No']}"

            lr_value = record["LR No"]
            if is_ok_marker(lr_value):
                record["LR No"] = "OK"
            elif include_blank_lr and not lr_value:
                record["LR No"] = "OK"
            else:
                report["rows_with_lr"] += 1
                continue     # LR number already received - not needed here

            data.append(record)
            kept_here += 1

        report["sheets"].append({
            "name": sheet_name, "used": True, "kept": kept_here,
            "header_row": header["row"] + 1,
        })

    report["columns_missing"] = [
        c for c in FINAL_COLUMNS if c not in report["columns_found"]
    ]
    report["rows_out"] = len(data)
    return data, report


# -------------------------------------------------------------- write out


def build_workbook(data, title="LR Pending"):
    wb = Workbook()
    sheet = wb.active
    sheet.title = title[:31]

    header_fill = PatternFill("solid", fgColor="15202B")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="C8D2DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sheet.append(FINAL_COLUMNS)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for record in data:
        sheet.append([record.get(column, "") for column in FINAL_COLUMNS])

    widths = [28, 34, 14, 10, 14, 9]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row,
                               max_col=len(FINAL_COLUMNS)):
        for cell in row:
            cell.border = border
        row[3].alignment = Alignment(horizontal="center")
        row[5].alignment = Alignment(horizontal="right")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:F{max(sheet.max_row, 1)}"

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
